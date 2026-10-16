"""Shared workbook styling: Arial, black text, grey shading only."""

from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .config import load_methodology

_fmt = load_methodology().get("formatting", {})
FONT_NAME = _fmt.get("font", "Arial")
NUM = _fmt.get("number_format", '#,##0.00;(#,##0.00);"-"')
DATE = _fmt.get("date_format", "dd mmm yyyy")

HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")
SUBTOTAL_FILL = PatternFill("solid", fgColor="F2F2F2")
_thin = Side(style="thin", color="808080")
BOX = Border(top=_thin, bottom=_thin, left=_thin, right=_thin)


def font(bold: bool = False, size: int = 10) -> Font:
    return Font(name=FONT_NAME, bold=bold, size=size)


def header(ws, row: int, labels: list[str], start_col: int = 1) -> None:
    for j, label in enumerate(labels):
        c = ws.cell(row, start_col + j, label)
        c.font = font(True)
        c.fill = HEADER_FILL
        c.border = BOX
        c.alignment = Alignment(wrap_text=True, vertical="center")
