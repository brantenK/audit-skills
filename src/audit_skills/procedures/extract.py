"""Client-facing ledger extract: the selected lines in the ledger's own columns, and a total."""

from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from ..ledger import GL_COLUMNS, GLLine
from ..xlstyle import BOX, DATE, NUM, font, header


def gl_extract(lines: list[GLLine], out: Path, entity: str, title: str, source_note: str) -> dict:
    """Write the extract. Returns line count and totals for checking."""
    lines = sorted(lines, key=lambda ln: (ln.account, ln.date or 0))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "GL extract"
    ws["A1"] = entity
    ws["A1"].font = font(True, 12)
    ws["A2"] = title
    ws["A2"].font = font(True)
    ws["A3"] = source_note
    ws["A3"].font = font()

    hr = 5
    header(ws, hr, GL_COLUMNS)
    r = hr + 1
    first = r
    for ln in lines:
        vals = [ln.account, ln.account_name, ln.date, ln.counterparty, ln.reference, ln.txn_type,
                ln.narrative, ln.debit, ln.credit, f"=H{r}-I{r}", ln.month]
        for j, v in enumerate(vals, 1):
            c = ws.cell(r, j, v)
            c.font = font()
            c.border = BOX
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(r, 3).number_format = DATE
        for col in (8, 9, 10):
            ws.cell(r, col).number_format = NUM
        r += 1
    last = r - 1
    tot = r + 1
    ws.cell(tot, 1, "Total").font = font(True)
    for col in (8, 9, 10):
        L = get_column_letter(col)
        c = ws.cell(tot, col, f"=SUM({L}{first}:{L}{last})" if lines else 0)
        c.font = font(True)
        c.number_format = NUM
    widths = [14, 34, 12, 36, 16, 16, 36, 14, 14, 14, 10]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    # New workbooks stay unfrozen; pane changes require an explicit user request.
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    debit = round(sum(ln.debit for ln in lines), 2)
    credit = round(sum(ln.credit for ln in lines), 2)
    return {"lines": len(lines), "debit": debit, "credit": credit, "net": round(debit - credit, 2),
            "path": str(out)}
