import json

import openpyxl
import pytest

from audit_skills import improve, install, stats
from audit_skills.config import Engagement, load_engagement, sample_requirement, save_engagement
from audit_skills.excel import approvals
from audit_skills.procedures import template


def test_scrub_removes_client_data():
    text = ("Client Acme Mining paid R 1,234,567.89 on account 4029741234; email a@b.co.za; "
            "file C:\\Users\\me\\Downloads\\Acme\\tb.xlsx; 63,257.66 to venue")
    out = improve.scrub(text, ["Acme Mining", "Acme"])
    assert "Acme" not in out
    assert "1,234,567" not in out and "63,257.66" not in out
    assert "4029741234" not in out and "a@b.co.za" not in out
    assert "C:\\Users" not in out


def test_propose_apply_and_install_merges_notes(tmp_path, monkeypatch):
    p = improve.propose("audit-conventions", "Always check trivial is R7,500 in every paper", [])
    assert "[amount]" in p.read_text() or "7,500" not in p.read_text()
    improve.apply(p.name)
    monkeypatch.chdir(tmp_path)
    done = install.install_skills("claude", project=True)
    assert len(done) == len(list(install.bundled_skills()))
    skill = (tmp_path / ".claude" / "skills" / "audit-conventions" / "SKILL.md").read_text()
    assert install.MARKER in skill and "Always check trivial" in skill
    # other agents' folders
    install.install_skills("codex", project=True)
    install.install_skills("opencode", project=True)
    assert (tmp_path / ".agents" / "skills" / "vouch-to-bank" / "SKILL.md").exists()
    assert (tmp_path / ".opencode" / "skills" / "tie-to-tb" / "SKILL.md").exists()


def test_every_skill_has_frontmatter():
    for sd in install.bundled_skills():
        text = (sd / "SKILL.md").read_text()
        assert text.startswith("---\n")
        head = text.split("---")[1]
        assert f"name: {sd.name}" in head
        assert "description:" in head


def test_mcp_config_written_for_codex_and_opencode(tmp_path):
    msg = install.install_mcp("codex")
    cfg = (tmp_path / "userhome" / ".codex" / "config.toml").read_text()
    assert "[mcp_servers.audit-skills]" in cfg and 'args = ["mcp"]' in cfg
    assert "already" in install.install_mcp("codex")
    install.install_mcp("opencode")
    data = json.loads((tmp_path / "userhome" / ".config" / "opencode" / "opencode.json").read_text())
    assert data["mcp"]["audit-skills"] == {"type": "local", "command": ["audit-skills", "mcp"], "enabled": True}


def test_engagement_roundtrip_and_sample_table(tmp_path):
    save_engagement(tmp_path, Engagement("Synthetic Ltd", "2026-06-30", 150000, 7500))
    eng = load_engagement(tmp_path)
    assert eng.individually_significant == 112500
    assert sample_requirement("Category C", "Medium") == {"number": 15, "coverage": 0.5}


def test_stats_off_by_default_and_minimal():
    stats.record("extract")
    assert stats.show() == []
    stats.set_enabled(True)
    stats.record("extract")
    entry = stats.show()[0]
    assert set(entry) == {"event", "version", "date"}


def test_approval_needs_a_terminal(tmp_path):
    wb = tmp_path / "paper.xlsx"
    with pytest.raises(PermissionError):
        approvals.grant(wb)  # pytest stdin is not a terminal
    approvals.grant(wb, interactive_check=False)
    assert approvals.is_approved(wb)
    approvals.revoke(wb)
    assert not approvals.is_approved(wb)


def test_template_inspect_and_roll(tmp_path):
    src = tmp_path / "py.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "2025"
    ws["A1"] = "Section: Prepaid"
    ws["B5"] = 2025
    ws["D5"] = 2024
    ws["A6"] = "Balances at 30 June 2025"
    ws["B7"] = "=SUM(1,2)"
    wb.create_sheet("Scratch")
    wb.save(src)
    info = template.inspect(src)
    assert info["cells_using_external_links"] == 0
    res = template.roll(src, tmp_path / "cy.xlsx", "2025", "2026")
    out = openpyxl.load_workbook(tmp_path / "cy.xlsx")
    assert "2026" in out.sheetnames and "Scratch" in out.sheetnames  # nothing deleted
    ws2 = out["2026"]
    assert ws2["B5"].value == 2026 and ws2["A6"].value == "Balances at 30 June 2026"
    assert ws2["B7"].value == "=SUM(1,2)"
    with pytest.raises(FileExistsError):
        template.roll(src, tmp_path / "cy.xlsx", "2025", "2026")
