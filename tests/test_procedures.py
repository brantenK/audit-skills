import datetime as dt

import openpyxl

from audit_skills.ledger import filter_accounts, read_balances, read_gl
from audit_skills.procedures import cutoff, extract, sample, tie, vouch
from audit_skills.ledger import _rows


def test_read_gl_and_balances(gl_file, tb_file):
    lines = read_gl(gl_file)
    assert len(lines) == 8
    assert lines[0].net == 50000
    tb = read_balances(tb_file)
    assert tb["9100/101"] == -140000


def test_extract_counts_and_totals(gl_file, tmp_path):
    lines = filter_accounts(read_gl(gl_file), ["8730/101", "8730/102", "8730/103"])
    res = extract.gl_extract(lines, tmp_path / "out.xlsx", "Synthetic Entity", "Prepaid extract", "Synthetic")
    assert res["lines"] == 4
    assert res["net"] == 81000
    ws = openpyxl.load_workbook(tmp_path / "out.xlsx").active
    assert ws["A5"].value == "Account_Code"
    # client extract has ledger columns only
    assert [ws.cell(5, c).value for c in range(1, 12)][-1] == "Month"
    assert ws.cell(5, 12).value is None


def test_rollforward_bridge_and_unposted_adjustment(gl_file, tb_file, py_tb_file):
    lines = read_gl(gl_file)
    rows = tie.rollforward(lines, read_balances(tb_file), ["9100/101", "9100/102"], read_balances(py_tb_file))
    r101 = next(r for r in rows if r.account == "9100/101")
    # closing -140,000; movement -115,000 (credits 120,000 less debit 5,000) => opening -25,000
    assert r101.movement == -115000
    assert r101.opening_per_ledger == -25000
    assert r101.closing_per_ledger == r101.closing_per_tb
    # prior-year audited TB shows nil: an auditor reclassification the client never posted
    assert r101.opening_difference == -25000
    b = tie.bridge(rows)
    assert b["agrees"]


def test_sample_plan_material_and_remainder():
    balances = {"A": 500000, "B": 200000, "C": 50000, "D": 30000, "E": 20000}
    p = sample.plan(balances, pm=150000, individually_significant=112500,
                    requirement={"number": 15, "coverage": 0.5})
    assert set(p.material) == {"A", "B"}
    assert p.remainder_total == 100000
    assert not p.remainder_is_material
    sample.select_remainder(p)
    assert p.coverage >= 0.5
    assert "satisfied" in p.sentence()


def test_sample_plan_without_risk_level():
    p = sample.plan({"A": 10}, 150000, 112500, None)
    assert p.sufficient is None
    assert "not set" in p.sentence()


def test_vouch_matches_payments_and_flags_invoices(gl_file, bank_file):
    lines = filter_accounts(read_gl(gl_file), ["8730/101", "8730/102", "8730/103"])
    matches = vouch.match(lines, vouch.read_bank(_rows(bank_file)), days=3)
    status = {(m.gl.account, m.gl.reference, m.gl.debit): m.status for m in matches}
    assert status[("8730/101", "pmt", 50000)] == "agreed"
    assert status[("8730/101", "SIV001", 20000)] == "not_a_bank_item"
    assert status[("8730/102", "pmt", 8000)] == "agreed"
    assert status[("8730/103", "pmt", 3000)] == "not_found"
    s = vouch.summary(matches)
    assert s["agreed"] == 2 and s["agreed_amount"] == 58000


def test_cutoff_flags():
    ye = dt.date(2026, 6, 30)
    rows = cutoff.assess([
        cutoff.Event("8730/101", "Expo 2027 (FY2028)", dt.date(2027, 4, 5)),
        cutoff.Event("8730/102", "Summit 2026", dt.date(2026, 7, 7)),
        cutoff.Event("8730/104", "Big Congress 2028", dt.date(2028, 9, 3), confirmed=False),
        cutoff.Event("8730/105", "Past event", dt.date(2026, 3, 1)),
    ], ye)
    by = {r["account"]: r for r in rows}
    assert by["8730/101"]["financial_year"] == 2027
    assert any("FY2028" in f for f in by["8730/101"]["flags"])
    assert by["8730/102"]["after_year_end"] is True and by["8730/102"]["flags"] == []
    assert any("non-current" in f for f in by["8730/104"]["flags"])
    assert any("not confirmed" in f for f in by["8730/104"]["flags"])
    assert any("released" in f for f in by["8730/105"]["flags"])
